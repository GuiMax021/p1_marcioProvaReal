"""
Gera data/acervo_consolidado.xlsx a partir de data/acervo.csv,
com formatacao profissional e uma aba de resumo com formulas.
"""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent.parent
CSV_PATH = BASE / "data" / "acervo.csv"
XLSX_PATH = BASE / "data" / "acervo_consolidado.xlsx"

FONT_NAME = "Arial"
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=11)
BODY_FONT = Font(name=FONT_NAME, size=11)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

CATEGORIAS = [
    ("000", "Generalidades e Informacao"),
    ("100", "Filosofia e Psicologia"),
    ("200", "Religiao e Teologia"),
    ("300", "Ciencias Sociais e Direito"),
    ("400", "Linguistica e Idiomas"),
    ("500", "Ciencias Puras (Exatas e Naturais)"),
    ("600", "Ciencias Aplicadas (Tecnologia)"),
    ("700", "Artes e Recreacao"),
    ("800", "Literatura"),
    ("900", "Historia e Geografia"),
]


def carregar_csv():
    with open(CSV_PATH, newline="", encoding="utf-8") as f:
        reader = csv.reader(f)
        header = next(reader)
        rows = list(reader)
    return header, rows


def montar_acervo(ws, header, rows):
    ws.title = "Acervo"
    ws.append(header)
    for row in rows:
        ws.append(row)

    for col_idx, _ in enumerate(header, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER

    last_row = ws.max_row
    last_col = len(header)
    for r in range(2, last_row + 1):
        for c in range(1, last_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center")

    widths = [8, 32, 22, 16, 26, 14, 8, 20, 18]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(last_col)}{last_row}"
    return last_row


def montar_resumo(ws, last_row_acervo):
    ws.title = "Resumo"
    ws.append(["Resumo do Acervo"])
    ws["A1"].font = Font(name=FONT_NAME, bold=True, size=14)
    ws.append([])

    # --- Por tipo de acervo ---
    ws.append(["Tipo de Acervo", "Quantidade"])
    for cell in ws[3]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center")

    tipos = ["Digital", "Fisico"]
    row = 4
    for tipo in tipos:
        ws.cell(row=row, column=1, value=tipo).font = BODY_FONT
        formula = f'=COUNTIF(Acervo!F2:F{last_row_acervo},"{tipo}")'
        c = ws.cell(row=row, column=2, value=formula)
        c.font = BODY_FONT
        ws.cell(row=row, column=1).border = BORDER
        c.border = BORDER
        row += 1

    total_row = row
    ws.cell(row=total_row, column=1, value="Total").font = Font(name=FONT_NAME, bold=True)
    total_formula = f"=SUM(B4:B{total_row - 1})"
    c = ws.cell(row=total_row, column=2, value=total_formula)
    c.font = Font(name=FONT_NAME, bold=True)

    row = total_row + 2

    # --- Por categoria (CDD) ---
    header_row = row
    ws.cell(row=header_row, column=1, value="Categoria (CDD)").font = HEADER_FONT
    ws.cell(row=header_row, column=2, value="Quantidade").font = HEADER_FONT
    for c in (1, 2):
        cell = ws.cell(row=header_row, column=c)
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center")
    row += 1

    first_cat_row = row
    for codigo, nome in CATEGORIAS:
        label = f"{codigo} - {nome}"
        ws.cell(row=row, column=1, value=label).font = BODY_FONT
        formula = f'=COUNTIF(Acervo!D2:D{last_row_acervo},"{codigo}")'
        c = ws.cell(row=row, column=2, value=formula)
        c.font = BODY_FONT
        ws.cell(row=row, column=1).border = BORDER
        c.border = BORDER
        row += 1
    last_cat_row = row - 1

    ws.cell(row=row, column=1, value="Total").font = Font(name=FONT_NAME, bold=True)
    total_cat_formula = f"=SUM(B{first_cat_row}:B{last_cat_row})"
    ws.cell(row=row, column=2, value=total_cat_formula).font = Font(name=FONT_NAME, bold=True)

    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 14


def main():
    header, rows = carregar_csv()

    wb = Workbook()
    ws_acervo = wb.active
    last_row = montar_acervo(ws_acervo, header, rows)

    ws_resumo = wb.create_sheet("Resumo")
    montar_resumo(ws_resumo, last_row)

    XLSX_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(XLSX_PATH)
    print(f"Planilha gerada em: {XLSX_PATH}")


if __name__ == "__main__":
    main()
